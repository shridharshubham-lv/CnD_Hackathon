import streamlit as st
import json
import io
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from app import defaults
from utils.stepper import render_stepper
from utils.sidebar import render_sidebar

for key, val in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = val

st.set_page_config(page_title="Knowledge Base & Export", page_icon="📚", layout="wide")

st.title("📚 Knowledge Base & Export")

render_sidebar(6)
render_stepper(6)

if not st.session_state.patterns:
    st.warning("Please complete the QA Report first — no patterns extracted yet.")
    st.page_link("pages/5_qa_report.py", label="← Go to QA Report", icon="⬅️")
    st.stop()

_eb = st.session_state.get("enriched_brief") or {}
_cname = _eb.get("campaignName", "Your Campaign")
st.markdown(
    f"""
    <div style="text-align:center;padding:1.5rem 1rem;margin-bottom:1rem;
                background:linear-gradient(135deg,#ecfdf5 0%,#eff6ff 100%);
                border:1.5px solid #86efac;border-radius:12px;">
        <div style="font-size:2.5rem;margin-bottom:0.3rem;">✅</div>
        <div style="font-size:1.4rem;font-weight:700;color:#166534;">Campaign Complete</div>
        <div style="font-size:1rem;color:#15803d;margin-top:0.2rem;">{_cname}</div>
        <div style="font-size:0.82rem;color:#64748b;margin-top:0.5rem;">
            Intelligence logged to memory · Patterns extracted · Ready for export
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

st.subheader("Patterns Learned from This Campaign")
for pattern in st.session_state.patterns:
    st.info(
        f"**Channel:** {pattern.get('channel', 'N/A')} | "
        f"**Error Type:** {pattern.get('errorType', 'N/A')}\n\n"
        f"**Lesson:** {pattern.get('lesson', '')}\n\n"
        f"**Trigger Condition:** {pattern.get('triggerCondition', '')}"
    )

st.divider()

st.subheader("How These Will Help Future Campaigns")
st.write("When a similar brief is analyzed in the future, you'll see:")
for pattern in st.session_state.patterns:
    st.warning(
        f"⚠️ Past campaigns with similar briefs had issues: "
        f"**{pattern.get('lesson', '')}** — "
        f"Triggered when: {pattern.get('triggerCondition', 'N/A')}"
    )

st.divider()

st.subheader("Export")
col1, col2 = st.columns(2)

with col1:
    export_data = {
        "campaign_id": st.session_state.campaign_id,
        "raw_brief": st.session_state.raw_brief,
        "enriched_brief": st.session_state.enriched_brief,
        "execution_plan": st.session_state.execution_plan,
        "channel_assets": st.session_state.get("channel_assets", {}),
        "qa_report": st.session_state.qa_report,
        "patterns": st.session_state.patterns,
        "conflicts": st.session_state.conflicts,
    }
    st.download_button(
        label="📄 Export Full Report (JSON)",
        data=json.dumps(export_data, indent=2, default=str),
        file_name=f"campaign_{st.session_state.campaign_id}_full_report.json",
        mime="application/json",
    )

with col2:
    # Build readable text summary
    lines = []
    lines.append("=" * 60)
    lines.append("CAMPAIGN QA SUMMARY")
    lines.append("=" * 60)
    lines.append(f"Campaign ID: {st.session_state.campaign_id}")
    lines.append(f"Brief: {st.session_state.raw_brief[:200]}...")
    lines.append("")

    qa = st.session_state.qa_report or []
    summary = st.session_state.get("qa_summary", {})
    lines.append(f"Total Issues: {summary.get('total', len(qa))}")
    lines.append(f"Critical: {summary.get('critical', 0)}")
    lines.append(f"Warning: {summary.get('warning', 0)}")
    lines.append(f"Advisory: {summary.get('advisory', 0)}")
    lines.append("")

    for issue in qa:
        lines.append(f"[{issue.get('severity', 'N/A')}] {issue.get('channel', '')} — {issue.get('dimension', '')}")
        lines.append(f"  Current: {issue.get('current_text', '')}")
        lines.append(f"  Required: {issue.get('brief_requirement', '')}")
        lines.append(f"  Fix: {issue.get('suggested_fix', '')}")
        lines.append(f"  Impact: {issue.get('business_impact', '')}")
        lines.append("")

    lines.append("PATTERNS LEARNED:")
    for p in st.session_state.patterns:
        lines.append(f"  - [{p.get('channel', '')}] {p.get('lesson', '')}")

    text_summary = "\n".join(lines)
    st.download_button(
        label="📝 Export QA Summary (Text)",
        data=text_summary,
        file_name=f"campaign_{st.session_state.campaign_id}_qa_summary.txt",
        mime="text/plain",
    )

# --- Channel Assets Export ---
st.divider()

st.subheader("📎 Export Channel Assets")

channel_assets = st.session_state.get("channel_assets", {})
assets_with_content = {ch: copy for ch, copy in channel_assets.items() if copy.strip()}

if not assets_with_content:
    st.info("No channel assets to export. Assets are captured in Stage 4.")
else:
    st.markdown(f"**{len(assets_with_content)} channel(s)** with content ready to export.")

    # Preview
    with st.expander("Preview assets", expanded=False):
        for ch, copy in assets_with_content.items():
            st.markdown(f"**{ch}**")
            st.code(copy, language=None)

    def _build_assets_text(assets: dict) -> str:
        lines = []
        for ch, copy in assets.items():
            lines.append("=" * 60)
            lines.append(f"CHANNEL: {ch}")
            lines.append("=" * 60)
            lines.append(copy)
            lines.append("")
        return "\n".join(lines)

    def _build_assets_docx(assets: dict) -> bytes:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_heading("Campaign Channel Assets", level=0)
        for ch, copy in assets.items():
            doc.add_heading(ch, level=1)
            para = doc.add_paragraph(copy)
            para.style.font.size = Pt(11)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def _build_assets_pdf(assets: dict) -> bytes:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 12, "Campaign Channel Assets", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(6)
        for ch, copy in assets.items():
            pdf.set_font("Helvetica", "B", 14)
            pdf.cell(0, 10, ch, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 11)
            # encode to latin-1 safe text for fpdf
            safe_copy = copy.encode("latin-1", "replace").decode("latin-1")
            pdf.multi_cell(0, 6, safe_copy)
            pdf.ln(8)
        return pdf.output()

    asset_col1, asset_col2, asset_col3 = st.columns(3)

    with asset_col1:
        st.download_button(
            label="📝 Download Assets (.txt)",
            data=_build_assets_text(assets_with_content),
            file_name=f"campaign_{st.session_state.campaign_id}_assets.txt",
            mime="text/plain",
        )

    with asset_col2:
        st.download_button(
            label="📄 Download Assets (.docx)",
            data=_build_assets_docx(assets_with_content),
            file_name=f"campaign_{st.session_state.campaign_id}_assets.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    with asset_col3:
        st.download_button(
            label="📕 Download Assets (.pdf)",
            data=_build_assets_pdf(assets_with_content),
            file_name=f"campaign_{st.session_state.campaign_id}_assets.pdf",
            mime="application/pdf",
        )

st.divider()

mem_nav_col1, mem_nav_col2 = st.columns(2)
with mem_nav_col1:
    if st.button("⬅️ Back to QA Report"):
        st.session_state.stage = 5
        st.switch_page("pages/5_qa_report.py")
with mem_nav_col2:
    if st.button("🔄 Start New Campaign"):
        from app import reset_campaign_state
        reset_campaign_state()
        st.switch_page("pages/1_brief_intake.py")
